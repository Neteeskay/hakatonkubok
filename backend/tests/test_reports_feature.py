from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models.enums import HelpCategory, ParticipationFormat, TaskType, UserRole
from app.schemas.reports import ParticipantReportRow, PlatformAnalyticsReport
from app.services.report_service import (
    _format_best_month_label,
    _format_volunteer_display_name,
    build_volunteer_year_statistics_pdf,
    build_analytics_csv,
    build_full_report_xlsx_bytes,
    build_participants_csv,
    build_participants_xlsx_bytes,
)


def make_participant_row() -> ParticipantReportRow:
    return ParticipantReportRow(
        volunteer_id=uuid4(),
        full_name="Ivan Petrov",
        email="volunteer@stoloto.local",
        city="Nizhny Novgorod",
        department="IT",
        position="Developer",
        registered_at=datetime.now(UTC),
        applications_count=2,
        completed_tasks_count=1,
        awarded_hours=Decimal("4.00"),
    )


def make_analytics() -> PlatformAnalyticsReport:
    return PlatformAnalyticsReport(
        volunteers_total=1,
        funds_total=1,
        funds_pending_review=0,
        funds_approved=1,
        tasks_total=2,
        tasks_pending_review=0,
        tasks_published=1,
        tasks_closed=1,
        applications_total=2,
        accepted_applications=1,
        completions_waiting_hours=1,
        awarded_hours_total=Decimal("4.00"),
    )


def test_participants_report_exports_csv_and_xlsx() -> None:
    rows = [make_participant_row()]

    csv_content = build_participants_csv(rows)
    xlsx_content = build_participants_xlsx_bytes(rows)

    assert "volunteer@stoloto.local" in csv_content
    assert xlsx_content.startswith(b"PK")
    workbook = load_workbook(BytesIO(xlsx_content))
    worksheet = workbook["Участники"]
    assert worksheet["A1"].value.startswith("Сформировано платформой")
    assert worksheet["A2"].value == "Отчет по участникам корпоративного волонтерства"
    assert worksheet["A3"].value == "ID волонтера"
    assert "A1:J1" in {str(cell_range) for cell_range in worksheet.merged_cells.ranges}


def test_full_report_export_contains_workbook_bytes() -> None:
    rows = [make_participant_row()]
    analytics = make_analytics()

    csv_content = build_analytics_csv(analytics)
    xlsx_content = build_full_report_xlsx_bytes(rows, analytics)

    assert "volunteers_total" in csv_content
    assert xlsx_content.startswith(b"PK")
    workbook = load_workbook(BytesIO(xlsx_content))
    assert "Участники" in workbook.sheetnames
    assert "Аналитика" in workbook.sheetnames
    assert workbook["Аналитика"]["A2"].value == "Сводная аналитика платформы"
    assert workbook["Аналитика"]["A4"].value == "Всего волонтеров"


def test_format_best_month_label_returns_russian_month_name() -> None:
    label = _format_best_month_label({8: Decimal("8.00"), 2: Decimal("2.00")})

    assert label == "Август"


def test_format_volunteer_display_name_prefers_full_name() -> None:
    volunteer = SimpleNamespace(full_name="Иван Петров", email="ivan@stoloto.local")

    assert _format_volunteer_display_name(volunteer) == "Иван Петров"


@pytest.mark.asyncio
async def test_volunteer_year_statistics_pdf_exports_pdf_bytes() -> None:
    class FakeResult:
        def all(self) -> list[SimpleNamespace]:
            return [
                SimpleNamespace(
                    title="Помощь на мероприятии",
                    fund_name="Фонд добрых дел",
                    category=HelpCategory.CHILDREN,
                    participation_format=ParticipationFormat.OFFLINE,
                    task_type=TaskType.REGULAR,
                    hours=Decimal("4.50"),
                    awarded_at=datetime(2026, 2, 12, 15, 30, tzinfo=UTC),
                )
            ]

    class FakeSession:
        async def execute(self, statement: object) -> FakeResult:
            return FakeResult()

    volunteer = SimpleNamespace(
        id=uuid4(),
        role=UserRole.VOLUNTEER,
        email="ivan.petrov@stoloto.local",
        full_name="Иван Петров",
        city="Нижний Новгород",
        department="IT",
        position="Backend developer",
    )

    pdf_content = await build_volunteer_year_statistics_pdf(
        FakeSession(),
        volunteer=volunteer,
        year=2026,
    )

    assert pdf_content.startswith(b"%PDF")
    assert len(pdf_content) > 10_000
