import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Iterable
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.domain import Fund, TaskApplication, User, VolunteerHourLedger, VolunteerTask
from app.models.enums import ApplicationStatus, FundStatus, TaskStatus, UserRole
from app.schemas.reports import ParticipantReportRow, PlatformAnalyticsReport

from collections import Counter, defaultdict
from pathlib import Path

from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.fonts import addMapping
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import registerFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PARTICIPANT_HEADERS = [
    "volunteer_id",
    "full_name",
    "email",
    "city",
    "department",
    "position",
    "registered_at",
    "applications_count",
    "completed_tasks_count",
    "awarded_hours",
]

ANALYTICS_HEADERS = ["metric", "value"]

PARTICIPANT_XLSX_HEADERS = [
    "ID волонтера",
    "ФИО",
    "Email",
    "Город",
    "Подразделение",
    "Должность",
    "Дата регистрации",
    "Откликов",
    "Выполнено задач",
    "Начислено часов",
]

ANALYTICS_LABELS = {
    "volunteers_total": "Всего волонтеров",
    "funds_total": "Всего фондов",
    "funds_pending_review": "Фондов на модерации",
    "funds_approved": "Одобренных фондов",
    "tasks_total": "Всего заданий",
    "tasks_pending_review": "Заданий на модерации",
    "tasks_published": "Опубликованных заданий",
    "tasks_closed": "Закрытых заданий",
    "applications_total": "Всего откликов",
    "accepted_applications": "Принятых откликов",
    "completions_waiting_hours": "Выполнений без начисленных часов",
    "awarded_hours_total": "Начислено часов всего",
}

ANALYTICS_XLSX_HEADERS = ["Показатель", "Значение"]

PARTICIPANT_COLUMN_WIDTHS = [38, 28, 32, 22, 24, 24, 22, 14, 18, 18]
ANALYTICS_COLUMN_WIDTHS = [42, 18, 18, 18]

THIN_BLACK_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
TITLE_FILL = PatternFill("solid", fgColor="F8F5FF")
HEADER_FILL = PatternFill("solid", fgColor="EDE7F6")


async def fetch_participant_report_rows(
    session: AsyncSession,
    *,
    limit: int | None,
    offset: int = 0,
) -> list[ParticipantReportRow]:
    application_counts = (
        select(
            TaskApplication.volunteer_id.label("volunteer_id"),
            func.count(TaskApplication.id).label("applications_count"),
        )
        .group_by(TaskApplication.volunteer_id)
        .subquery()
    )
    hour_totals = (
        select(
            VolunteerHourLedger.volunteer_id.label("volunteer_id"),
            func.count(VolunteerHourLedger.id).label("completed_tasks_count"),
            func.coalesce(func.sum(VolunteerHourLedger.hours), 0).label("awarded_hours"),
        )
        .group_by(VolunteerHourLedger.volunteer_id)
        .subquery()
    )

    statement = (
        select(
            User.id.label("volunteer_id"),
            User.full_name,
            User.email,
            User.city,
            User.department,
            User.position,
            User.created_at.label("registered_at"),
            func.coalesce(application_counts.c.applications_count, 0).label(
                "applications_count"
            ),
            func.coalesce(hour_totals.c.completed_tasks_count, 0).label(
                "completed_tasks_count"
            ),
            func.coalesce(hour_totals.c.awarded_hours, 0).label("awarded_hours"),
        )
        .outerjoin(application_counts, application_counts.c.volunteer_id == User.id)
        .outerjoin(hour_totals, hour_totals.c.volunteer_id == User.id)
        .where(User.role == UserRole.VOLUNTEER)
        .order_by(User.created_at.desc())
        .offset(offset)
    )
    if limit is not None:
        statement = statement.limit(limit)

    result = await session.execute(statement)
    return [ParticipantReportRow(**row._mapping) for row in result.all()]


async def fetch_platform_analytics(session: AsyncSession) -> PlatformAnalyticsReport:
    volunteers_total = await _count(session, select(func.count(User.id)).where(User.role == UserRole.VOLUNTEER))
    funds_total = await _count(session, select(func.count(Fund.id)))
    funds_pending_review = await _count(
        session,
        select(func.count(Fund.id)).where(Fund.status == FundStatus.PENDING_REVIEW),
    )
    funds_approved = await _count(
        session,
        select(func.count(Fund.id)).where(Fund.status == FundStatus.APPROVED),
    )
    tasks_total = await _count(session, select(func.count(VolunteerTask.id)))
    tasks_pending_review = await _count(
        session,
        select(func.count(VolunteerTask.id)).where(VolunteerTask.status == TaskStatus.PENDING_REVIEW),
    )
    tasks_published = await _count(
        session,
        select(func.count(VolunteerTask.id)).where(VolunteerTask.status == TaskStatus.PUBLISHED),
    )
    tasks_closed = await _count(
        session,
        select(func.count(VolunteerTask.id)).where(VolunteerTask.status == TaskStatus.CLOSED),
    )
    applications_total = await _count(session, select(func.count(TaskApplication.id)))
    accepted_applications = await _count(
        session,
        select(func.count(TaskApplication.id)).where(TaskApplication.status == ApplicationStatus.ACCEPTED),
    )
    completions_waiting_hours = await _count(
        session,
        select(func.count(TaskApplication.id))
        .where(TaskApplication.status == ApplicationStatus.COMPLETION_CONFIRMED)
        .where(
            ~select(VolunteerHourLedger.id)
            .where(VolunteerHourLedger.application_id == TaskApplication.id)
            .exists()
        ),
    )
    awarded_hours_total = await session.scalar(
        select(func.coalesce(func.sum(VolunteerHourLedger.hours), 0))
    )

    return PlatformAnalyticsReport(
        volunteers_total=volunteers_total,
        funds_total=funds_total,
        funds_pending_review=funds_pending_review,
        funds_approved=funds_approved,
        tasks_total=tasks_total,
        tasks_pending_review=tasks_pending_review,
        tasks_published=tasks_published,
        tasks_closed=tasks_closed,
        applications_total=applications_total,
        accepted_applications=accepted_applications,
        completions_waiting_hours=completions_waiting_hours,
        awarded_hours_total=awarded_hours_total or 0,
    )


async def _count(session: AsyncSession, statement) -> int:
    return int(await session.scalar(statement) or 0)


def build_participants_csv(rows: Iterable[ParticipantReportRow]) -> str:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(PARTICIPANT_HEADERS)
    for row in rows:
        writer.writerow(_participant_row_values(row))
    return output.getvalue()


def build_participants_xlsx_bytes(rows: Iterable[ParticipantReportRow]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Участники"
    _apply_report_title(
        worksheet,
        title="Отчет по участникам корпоративного волонтерства",
        title_columns=len(PARTICIPANT_XLSX_HEADERS),
    )
    worksheet.append(PARTICIPANT_XLSX_HEADERS)
    for row in rows:
        worksheet.append(_participant_row_values(row))
    _style_table_sheet(
        worksheet,
        header_row=3,
        data_start_row=4,
        data_columns=len(PARTICIPANT_XLSX_HEADERS),
        column_widths=PARTICIPANT_COLUMN_WIDTHS,
    )
    return _workbook_to_bytes(workbook)


def build_analytics_csv(report: PlatformAnalyticsReport) -> str:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(ANALYTICS_HEADERS)
    for metric, value in report.model_dump().items():
        writer.writerow([metric, _export_value(value)])
    return output.getvalue()


def build_analytics_xlsx_bytes(report: PlatformAnalyticsReport) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Аналитика"
    _apply_report_title(
        worksheet,
        title="Сводная аналитика платформы",
        title_columns=len(ANALYTICS_COLUMN_WIDTHS),
    )
    worksheet.append(ANALYTICS_XLSX_HEADERS)
    for metric, value in report.model_dump().items():
        worksheet.append([ANALYTICS_LABELS.get(metric, metric), _export_value(value)])
    _style_table_sheet(
        worksheet,
        header_row=3,
        data_start_row=4,
        data_columns=len(ANALYTICS_XLSX_HEADERS),
        column_widths=ANALYTICS_COLUMN_WIDTHS,
    )
    return _workbook_to_bytes(workbook)


def build_full_report_xlsx_bytes(
    participants: Iterable[ParticipantReportRow],
    analytics: PlatformAnalyticsReport,
) -> bytes:
    workbook = Workbook()
    participants_sheet = workbook.active
    participants_sheet.title = "Участники"
    _apply_report_title(
        participants_sheet,
        title="Отчет по участникам корпоративного волонтерства",
        title_columns=len(PARTICIPANT_XLSX_HEADERS),
    )
    participants_sheet.append(PARTICIPANT_XLSX_HEADERS)
    for row in participants:
        participants_sheet.append(_participant_row_values(row))
    _style_table_sheet(
        participants_sheet,
        header_row=3,
        data_start_row=4,
        data_columns=len(PARTICIPANT_XLSX_HEADERS),
        column_widths=PARTICIPANT_COLUMN_WIDTHS,
    )

    analytics_sheet = workbook.create_sheet("Аналитика")
    _apply_report_title(
        analytics_sheet,
        title="Сводная аналитика платформы",
        title_columns=len(ANALYTICS_COLUMN_WIDTHS),
    )
    analytics_sheet.append(ANALYTICS_XLSX_HEADERS)
    for metric, value in analytics.model_dump().items():
        analytics_sheet.append([ANALYTICS_LABELS.get(metric, metric), _export_value(value)])
    _style_table_sheet(
        analytics_sheet,
        header_row=3,
        data_start_row=4,
        data_columns=len(ANALYTICS_XLSX_HEADERS),
        column_widths=ANALYTICS_COLUMN_WIDTHS,
    )

    return _workbook_to_bytes(workbook)


def _workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _apply_report_title(worksheet, *, title: str, title_columns: int) -> None:
    generated_at = datetime.now().strftime("%d.%m.%Y %H:%M")

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_columns)
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=title_columns)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"

    service_cell = worksheet.cell(row=1, column=1)
    service_cell.value = (
        "Сформировано платформой «Помогать просто»\n"
        f"Дата формирования: {generated_at}"
    )
    service_cell.font = Font(name="Times New Roman", size=12, bold=True)
    service_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    service_cell.fill = TITLE_FILL

    title_cell = worksheet.cell(row=2, column=1)
    title_cell.value = title
    title_cell.font = Font(name="Times New Roman", size=14, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_cell.fill = TITLE_FILL

    worksheet.row_dimensions[1].height = 54
    worksheet.row_dimensions[2].height = 48
    worksheet.row_dimensions[3].height = 42


def _style_table_sheet(
    worksheet,
    *,
    header_row: int,
    data_start_row: int,
    data_columns: int,
    column_widths: list[int],
) -> None:
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for row in worksheet.iter_rows(
        min_row=header_row,
        max_row=max(worksheet.max_row, header_row),
        min_col=1,
        max_col=data_columns,
    ):
        for cell in row:
            cell.border = THIN_BLACK_BORDER
            cell.font = Font(name="Times New Roman", size=12, bold=cell.row == header_row)
            cell.alignment = Alignment(
                horizontal="center" if cell.row == header_row else "left",
                vertical="center",
                wrap_text=True,
            )
            if cell.row == header_row:
                cell.fill = HEADER_FILL

    for row_number in range(data_start_row, worksheet.max_row + 1):
        worksheet.row_dimensions[row_number].height = 30


def _participant_row_values(row: ParticipantReportRow) -> list[object]:
    return [_export_value(getattr(row, field)) for field in PARTICIPANT_HEADERS]


def _export_value(value: object) -> object:
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


BRAND_YELLOW = colors.HexColor("#FFE300")
BRAND_BLACK = colors.HexColor("#000000")
PDF_SOFT_YELLOW = colors.HexColor("#FFF9CC")
PDF_LIGHT_GRAY = colors.HexColor("#F5F5F5")
PDF_MID_GRAY = colors.HexColor("#DADADA")
PDF_TEXT_GRAY = colors.HexColor("#555555")

PDF_PAGE_MARGIN_X = 16 * mm
PDF_PAGE_MARGIN_TOP = 14 * mm
PDF_PAGE_MARGIN_BOTTOM = 12 * mm
PDF_CONTENT_WIDTH = 210 * mm - (2 * PDF_PAGE_MARGIN_X)

MONTH_NAMES_RU = (
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
)

CATEGORY_LABELS = {
    "children": "Дети",
    "elderly": "Пожилые",
    "disability": "Люди с ОВЗ",
    "ecology": "Экология",
}

FONT_DIR = Path(__file__).resolve().parents[1] / "static" / "fonts"
MONTSERRAT_REGULAR_PATH = FONT_DIR / "Montserrat-Regular.ttf"
MONTSERRAT_BOLD_PATH = FONT_DIR / "Montserrat-Bold.ttf"


def register_montserrat_fonts() -> None:
    if not MONTSERRAT_REGULAR_PATH.exists() or not MONTSERRAT_BOLD_PATH.exists():
        raise RuntimeError(
            "Montserrat fonts are required. Put Montserrat-Regular.ttf and "
            "Montserrat-Bold.ttf into app/static/fonts/"
        )

    registerFont(TTFont("Montserrat", str(MONTSERRAT_REGULAR_PATH)))
    registerFont(TTFont("Montserrat-Bold", str(MONTSERRAT_BOLD_PATH)))
    addMapping("Montserrat", 0, 0, "Montserrat")
    addMapping("Montserrat", 1, 0, "Montserrat-Bold")


async def _build_volunteer_year_statistics_pdf_legacy(
    session: AsyncSession,
    *,
    volunteer: User,
    year: int,
) -> bytes:
    register_montserrat_fonts()

    result = await session.execute(
        select(
            VolunteerTask.title,
            Fund.name.label("fund_name"),
            VolunteerTask.category,
            VolunteerTask.participation_format,
            VolunteerTask.task_type,
            VolunteerHourLedger.hours,
            VolunteerHourLedger.awarded_at,
        )
        .join(TaskApplication, TaskApplication.id == VolunteerHourLedger.application_id)
        .join(VolunteerTask, VolunteerTask.id == VolunteerHourLedger.task_id)
        .join(Fund, Fund.id == VolunteerTask.fund_id)
        .where(VolunteerHourLedger.volunteer_id == volunteer.id)
        .where(func.extract("year", VolunteerHourLedger.awarded_at) == year)
        .order_by(VolunteerHourLedger.awarded_at.asc())
    )

    rows = result.all()

    total_hours = sum(Decimal(row.hours) for row in rows)
    completed_tasks_count = len(rows)

    months_hours: dict[int, Decimal] = defaultdict(lambda: Decimal("0"))
    category_counter: Counter[str] = Counter()

    for row in rows:
        months_hours[row.awarded_at.month] += Decimal(row.hours)
        category_counter[str(row.category.value)] += 1

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"volunteer_statistics_{year}",
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="MontserratTitle",
            fontName="Montserrat-Bold",
            fontSize=20,
            leading=24,
            textColor=BRAND_BLACK,
            alignment=TA_CENTER,
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratHeading",
            fontName="Montserrat-Bold",
            fontSize=13,
            leading=16,
            textColor=BRAND_BLACK,
            spaceBefore=12,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratText",
            fontName="Montserrat",
            fontSize=10,
            leading=13,
            textColor=BRAND_BLACK,
        )
    )

    story = []

    story.append(Paragraph("Личная статистика волонтёра", styles["MontserratTitle"]))
    story.append(Paragraph(f"Отчётный год: {year}", styles["MontserratText"]))
    story.append(Spacer(1, 8))

    summary_table = Table(
        [
            ["Волонтёр", volunteer.full_name or volunteer.email],
            ["Email", volunteer.email],
            ["Город", volunteer.city or "Не указан"],
            ["Подразделение", volunteer.department or "Не указано"],
            ["Выполнено заданий", str(completed_tasks_count)],
            ["Начислено часов", str(total_hours)],
        ],
        colWidths=[55 * mm, 105 * mm],
    )
    summary_table.setStyle(_pdf_table_style())
    story.append(summary_table)

    if rows:
        story.append(Paragraph("Часы по месяцам", styles["MontserratHeading"]))
        story.append(_build_month_hours_chart(months_hours))

        story.append(Paragraph("Категории помощи", styles["MontserratHeading"]))
        category_data = [["Категория", "Количество выполненных заданий"]]
        for category, count in category_counter.most_common():
            category_data.append([category, str(count)])

        category_table = Table(category_data, colWidths=[90 * mm, 70 * mm])
        category_table.setStyle(_pdf_table_style())
        story.append(category_table)

        story.append(Paragraph("Выполненные задания", styles["MontserratHeading"]))

        task_table_data = [["Дата", "Задание", "Фонд", "Часы"]]
        for row in rows:
            task_table_data.append(
                [
                    row.awarded_at.strftime("%d.%m.%Y"),
                    row.title,
                    row.fund_name,
                    str(row.hours),
                ]
            )

        task_table = Table(
            task_table_data,
            colWidths=[25 * mm, 65 * mm, 50 * mm, 20 * mm],
            repeatRows=1,
        )
        task_table.setStyle(_pdf_table_style())
        story.append(task_table)
    else:
        story.append(Spacer(1, 10))
        story.append(
            Paragraph(
                "За выбранный год подтверждённых выполнений и начисленных часов пока нет.",
                styles["MontserratText"],
            )
        )

    doc.build(story)
    return buffer.getvalue()


def _pdf_table_style() -> TableStyle:
    return TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), "Montserrat"),
            ("FONTNAME", (0, 0), (-1, 0), "Montserrat-Bold"),
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_YELLOW),
            ("TEXTCOLOR", (0, 0), (-1, -1), BRAND_BLACK),
            ("GRID", (0, 0), (-1, -1), 0.6, BRAND_BLACK),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )


async def build_volunteer_year_statistics_pdf(
    session: AsyncSession,
    *,
    volunteer: User,
    year: int,
) -> bytes:
    register_montserrat_fonts()

    result = await session.execute(
        select(
            VolunteerTask.title,
            Fund.name.label("fund_name"),
            VolunteerTask.category,
            VolunteerTask.participation_format,
            VolunteerTask.task_type,
            VolunteerHourLedger.hours,
            VolunteerHourLedger.awarded_at,
        )
        .join(TaskApplication, TaskApplication.id == VolunteerHourLedger.application_id)
        .join(VolunteerTask, VolunteerTask.id == VolunteerHourLedger.task_id)
        .join(Fund, Fund.id == VolunteerTask.fund_id)
        .where(VolunteerHourLedger.volunteer_id == volunteer.id)
        .where(func.extract("year", VolunteerHourLedger.awarded_at) == year)
        .order_by(VolunteerHourLedger.awarded_at.asc())
    )

    rows = result.all()
    total_hours = sum(Decimal(row.hours) for row in rows)
    completed_tasks_count = len(rows)

    months_hours: dict[int, Decimal] = defaultdict(lambda: Decimal("0"))
    category_counter: Counter[str] = Counter()

    for row in rows:
        months_hours[row.awarded_at.month] += Decimal(row.hours)
        category_counter[str(row.category.value)] += 1

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=PDF_PAGE_MARGIN_X,
        leftMargin=PDF_PAGE_MARGIN_X,
        topMargin=PDF_PAGE_MARGIN_TOP,
        bottomMargin=PDF_PAGE_MARGIN_BOTTOM,
        title=f"volunteer_statistics_{year}",
    )

    styles = _pdf_styles()
    volunteer_name = _format_volunteer_display_name(volunteer)
    best_month_label = _format_best_month_label(months_hours)

    story = [
        _build_pdf_hero(
            styles,
            volunteer_name=volunteer_name,
            year=year,
            total_hours=total_hours,
        ),
        Spacer(1, 8),
        _build_pdf_kpi_cards(
            completed_tasks_count=completed_tasks_count,
            total_hours=total_hours,
            best_month_label=best_month_label,
        ),
        Spacer(1, 8),
    ]

    if rows:
        story.extend(
            [
                _section_title("Динамика часов", styles),
                Spacer(1, 4),
                _build_month_hours_chart(months_hours),
                Spacer(1, 8),
                _section_title("Категории помощи", styles),
                Spacer(1, 4),
                _build_category_table(category_counter),
                Spacer(1, 8),
                _section_title("Выполненные задания", styles),
                Spacer(1, 4),
                _build_task_table(rows, styles),
            ]
        )
    else:
        story.append(
            _build_empty_state(
                "За выбранный год подтвержденных выполнений и начисленных часов пока нет.",
                styles,
            )
        )

    doc.build(story, onFirstPage=_draw_pdf_page, onLaterPages=_draw_pdf_page)
    return buffer.getvalue()


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="MontserratHeroTitle",
            fontName="Montserrat-Bold",
            fontSize=18,
            leading=21,
            textColor=BRAND_BLACK,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratHeroSubtitle",
            fontName="Montserrat-Bold",
            fontSize=12,
            leading=14,
            textColor=BRAND_BLACK,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratHeroMeta",
            fontName="Montserrat",
            fontSize=9,
            leading=12,
            textColor=PDF_TEXT_GRAY,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratSection",
            fontName="Montserrat-Bold",
            fontSize=10,
            leading=12,
            textColor=BRAND_BLACK,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratText",
            fontName="Montserrat",
            fontSize=8,
            leading=10,
            textColor=BRAND_BLACK,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratMuted",
            fontName="Montserrat",
            fontSize=7,
            leading=9,
            textColor=PDF_TEXT_GRAY,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MontserratTableCell",
            fontName="Montserrat",
            fontSize=7,
            leading=8.5,
            textColor=BRAND_BLACK,
        )
    )
    return styles


def _format_volunteer_display_name(volunteer: User) -> str:
    full_name = (volunteer.full_name or "").strip()
    return full_name or "Волонтёр"


def _format_best_month_label(months_hours: dict[int, Decimal]) -> str:
    if not months_hours:
        return "—"

    max_hours = max(months_hours.values())
    if max_hours <= 0:
        return "—"

    best_month = min(
        month for month, hours in months_hours.items() if hours == max_hours
    )
    return MONTH_NAMES_RU[best_month - 1]


def _draw_pdf_page(canvas, doc) -> None:
    width, height = A4
    canvas.saveState()
    canvas.setFillColor(BRAND_YELLOW)
    canvas.rect(0, height - 4 * mm, width, 4 * mm, stroke=0, fill=1)
    canvas.setFillColor(PDF_TEXT_GRAY)
    canvas.setFont("Montserrat", 7)
    canvas.drawRightString(
        width - PDF_PAGE_MARGIN_X,
        PDF_PAGE_MARGIN_BOTTOM - 2 * mm,
        f"Страница {doc.page}",
    )
    canvas.restoreState()


def _build_pdf_hero(
    styles,
    *,
    volunteer_name: str,
    year: int,
    total_hours: Decimal,
) -> Table:
    left = [
        Paragraph("Личная статистика волонтёра", styles["MontserratHeroTitle"]),
        Paragraph(volunteer_name, styles["MontserratHeroSubtitle"]),
        Paragraph(f"Отчётный год {year}", styles["MontserratHeroMeta"]),
    ]
    right = [
        Paragraph("Всего часов", styles["MontserratMuted"]),
        Paragraph(f"{total_hours}", styles["MontserratHeroTitle"]),
    ]
    hero_value_width = 48 * mm
    table = Table(
        [[left, right]],
        colWidths=[PDF_CONTENT_WIDTH - hero_value_width, hero_value_width],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_LIGHT_GRAY),
                ("BACKGROUND", (1, 0), (1, 0), BRAND_YELLOW),
                ("BOX", (0, 0), (-1, -1), 0.8, BRAND_BLACK),
                ("LINEBEFORE", (1, 0), (1, 0), 0.8, BRAND_BLACK),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (0, 0), 14),
                ("RIGHTPADDING", (0, 0), (0, 0), 10),
                ("LEFTPADDING", (1, 0), (1, 0), 10),
                ("RIGHTPADDING", (1, 0), (1, 0), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    return table


def _build_pdf_kpi_cards(
    *,
    completed_tasks_count: int,
    total_hours: Decimal,
    best_month_label: str,
) -> Table:
    kpi_width = PDF_CONTENT_WIDTH / 3
    data = [
        ["Выполнено задач", "Начислено часов", "Лучший месяц"],
        [str(completed_tasks_count), str(total_hours), best_month_label],
    ]
    table = Table(
        data,
        colWidths=[kpi_width, kpi_width, kpi_width],
        rowHeights=[9 * mm, 13 * mm],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_LIGHT_GRAY),
                ("BACKGROUND", (0, 1), (-1, 1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.7, PDF_MID_GRAY),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, PDF_MID_GRAY),
                ("FONTNAME", (0, 0), (-1, 0), "Montserrat"),
                ("FONTNAME", (0, 1), (-1, 1), "Montserrat-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("FONTSIZE", (0, 1), (1, 1), 15),
                ("FONTSIZE", (2, 1), (2, 1), 12),
                ("TEXTCOLOR", (0, 0), (-1, 0), PDF_TEXT_GRAY),
                ("TEXTCOLOR", (0, 1), (-1, 1), BRAND_BLACK),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _section_title(title: str, styles) -> Table:
    table = Table(
        [[Paragraph(title, styles["MontserratSection"])]],
        colWidths=[PDF_CONTENT_WIDTH],
        rowHeights=[8 * mm],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), PDF_SOFT_YELLOW),
                ("LINEBEFORE", (0, 0), (0, 0), 4, BRAND_YELLOW),
                ("BOTTOMPADDING", (0, 0), (0, 0), 3),
                ("TOPPADDING", (0, 0), (0, 0), 3),
                ("LEFTPADDING", (0, 0), (0, 0), 8),
            ]
        )
    )
    return table


def _build_category_table(category_counter: Counter[str]) -> Table:
    data = [["Категория", "Выполнено задач"]]
    for category, count in category_counter.most_common():
        data.append([CATEGORY_LABELS.get(category, category), str(count)])

    table = Table(data, colWidths=[PDF_CONTENT_WIDTH - 58 * mm, 58 * mm])
    table.setStyle(_pdf_modern_table_style())
    return table


def _build_task_table(rows, styles) -> Table:
    task_table_data = [["Дата", "Задание", "Фонд", "Часы"]]
    for row in rows:
        task_table_data.append(
            [
                row.awarded_at.strftime("%d.%m.%Y"),
                Paragraph(str(row.title), styles["MontserratTableCell"]),
                Paragraph(str(row.fund_name), styles["MontserratTableCell"]),
                str(row.hours),
            ]
        )

    table = Table(
        task_table_data,
        colWidths=[21 * mm, 74 * mm, 58 * mm, 21 * mm],
        repeatRows=1,
    )
    table.setStyle(_pdf_modern_table_style())
    return table


def _build_empty_state(message: str, styles) -> Table:
    table = Table(
        [[Paragraph(message, styles["MontserratText"])]],
        colWidths=[PDF_CONTENT_WIDTH],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), PDF_LIGHT_GRAY),
                ("BOX", (0, 0), (0, 0), 0.7, PDF_MID_GRAY),
                ("LEFTPADDING", (0, 0), (0, 0), 12),
                ("RIGHTPADDING", (0, 0), (0, 0), 12),
                ("TOPPADDING", (0, 0), (0, 0), 12),
                ("BOTTOMPADDING", (0, 0), (0, 0), 12),
            ]
        )
    )
    return table


def _pdf_modern_table_style() -> TableStyle:
    return TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), "Montserrat"),
            ("FONTNAME", (0, 0), (-1, 0), "Montserrat-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 7.5),
            ("FONTSIZE", (0, 1), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_BLACK),
            ("TEXTCOLOR", (0, 0), (-1, 0), BRAND_YELLOW),
            ("TEXTCOLOR", (0, 1), (-1, -1), BRAND_BLACK),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PDF_LIGHT_GRAY]),
            ("GRID", (0, 0), (-1, -1), 0.35, PDF_MID_GRAY),
            ("BOX", (0, 0), (-1, -1), 0.8, BRAND_BLACK),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("ALIGN", (-1, 1), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
    )


def _build_month_hours_chart_legacy(months_hours: dict[int, Decimal]) -> Drawing:
    month_labels = [
        "Янв", "Фев", "Мар", "Апр", "Май", "Июн",
        "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек",
    ]

    values = [float(months_hours.get(month, Decimal("0"))) for month in range(1, 13)]

    drawing = Drawing(460, 220)

    chart = VerticalBarChart()
    chart.x = 35
    chart.y = 35
    chart.height = 150
    chart.width = 390
    chart.data = [values]
    chart.categoryAxis.categoryNames = month_labels
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = max(values) + 1 if max(values) > 0 else 1
    chart.valueAxis.valueStep = max(1, int(chart.valueAxis.valueMax / 5))
    chart.bars[0].fillColor = BRAND_YELLOW
    chart.bars[0].strokeColor = BRAND_BLACK
    chart.categoryAxis.labels.fontName = "Montserrat"
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.labels.fontName = "Montserrat"
    chart.valueAxis.labels.fontSize = 7

    drawing.add(chart)
    drawing.add(
        String(
            35,
            200,
            "Начисленные волонтёрские часы по месяцам",
            fontName="Montserrat-Bold",
            fontSize=10,
            fillColor=BRAND_BLACK,
        )
    )

    return drawing


def _build_month_hours_chart(months_hours: dict[int, Decimal]) -> Drawing:
    month_labels = [
        "Янв",
        "Фев",
        "Мар",
        "Апр",
        "Май",
        "Июн",
        "Июл",
        "Авг",
        "Сен",
        "Окт",
        "Ноя",
        "Дек",
    ]
    values = [float(months_hours.get(month, Decimal("0"))) for month in range(1, 13)]
    max_value = max(values) if values else 0

    chart_width = float(PDF_CONTENT_WIDTH)
    drawing = Drawing(chart_width, 145)
    drawing.add(
        Rect(0, 0, chart_width, 145, fillColor=PDF_LIGHT_GRAY, strokeColor=PDF_MID_GRAY)
    )
    drawing.add(Rect(0, 0, chart_width, 6, fillColor=BRAND_YELLOW, strokeColor=BRAND_YELLOW))
    drawing.add(
        String(
            22,
            123,
            "Начисленные часы по месяцам",
            fontName="Montserrat-Bold",
            fontSize=11,
            fillColor=BRAND_BLACK,
        )
    )

    chart = VerticalBarChart()
    chart.x = 30
    chart.y = 24
    chart.height = 88
    chart.width = chart_width - 60
    chart.data = [values]
    chart.categoryAxis.categoryNames = month_labels
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = max_value + 2 if max_value > 0 else 1
    chart.valueAxis.valueStep = max(1, int(chart.valueAxis.valueMax / 4))
    chart.bars[0].fillColor = BRAND_YELLOW
    chart.bars[0].strokeColor = BRAND_BLACK
    chart.barSpacing = 3
    chart.groupSpacing = 8
    chart.categoryAxis.strokeColor = PDF_TEXT_GRAY
    chart.valueAxis.strokeColor = PDF_TEXT_GRAY
    chart.categoryAxis.labels.fontName = "Montserrat"
    chart.categoryAxis.labels.fontSize = 7
    chart.categoryAxis.labels.fillColor = BRAND_BLACK
    chart.valueAxis.labels.fontName = "Montserrat"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.labels.fillColor = PDF_TEXT_GRAY

    drawing.add(chart)
    return drawing
